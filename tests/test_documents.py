from pathlib import Path

import fitz
from docx import Document

from procurement_bot.documents import extract_docx, extract_pdf, joined_text


def test_extract_pdf_preserves_page_locator(tmp_path: Path):
    path = tmp_path / "request.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Cable 3x2.5 - 200 m")
    document.save(path)
    document.close()

    segments = extract_pdf(path)

    assert segments[0].locator == "page:1"
    assert "Cable" in segments[0].text
    assert "[page:1]" in joined_text(segments)


def test_extract_docx_reads_paragraphs_and_table_rows(tmp_path: Path):
    path = tmp_path / "request.docx"
    document = Document()
    document.add_paragraph("Город: Алматы")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Марля"
    table.cell(0, 1).text = "25 м"
    document.save(path)

    segments = extract_docx(path)

    assert [segment.kind for segment in segments] == ["docx_paragraph", "docx_table_row"]
    assert "Марля | 25 м" in segments[1].text
