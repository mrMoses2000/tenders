"""Synthetic native queue integration. No live browser, Telegram, WAHA or LLM.

Offers must originate in the native browser validation/persistence path. Keep this
fixture independent of owner-specific test modules, whose portability is reviewed
separately. Only per-test schemas on the exact disposable target are writable.
"""

from __future__ import annotations

import hashlib
import hmac
import json
import os
import time
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import urlencode
from uuid import uuid4

import asyncpg
import fitz
import pytest
from aiohttp.test_utils import TestClient, TestServer
from docx import Document

from procurement_bot.approvals import approve_request
from procurement_bot.config import Settings
from procurement_bot.db import _init_connection, run_migrations
from procurement_bot.intake import IntakeService, ParsedProcurementRequest, ProcurementRequestPatch
from procurement_bot.mini_app import create_mini_app
from procurement_bot.outbox import WhatsAppOutboxWorker
from procurement_bot.owner_autonomy import ResearchAutomationPolicy, replace_research_profile
from procurement_bot.parallel_research import ParallelResearchRunExecutor
from procurement_bot.providers.browser import BROWSER_SPEC_SHA256, SupplierBrowserResult
from procurement_bot.providers.waha import WahaAccountIdentity, WahaContactCheck, WahaSendResult
from procurement_bot.queue import IncomingAttachment, persist_telegram_message
from procurement_bot.quote_snapshots import get_current_quote_snapshot
from procurement_bot.research import (
    ContactKind,
    EvidenceClaim,
    EvidenceType,
    ResearchSource,
    SupplierCandidate,
    SupplierContact,
    SupplierEvidence,
)
from procurement_bot.research_execution import _request_hash
from procurement_bot.research_workflow import start_approved_research
from procurement_bot.storage import LocalContentAddressedStorage
from procurement_bot.waha_webhook_http import WahaWebhookHttpIngress, WahaWebhookHttpRequest
from procurement_bot.whatsapp_account import (
    configure_account_budget,
    ensure_account,
    observe_provider_state,
    pin_whatsapp_account_identity,
    set_account_hold,
)
from procurement_bot.whatsapp_approval_workflow import start_approved_whatsapp_send
from procurement_bot.worker import JobWorker
from scripts.postgres_test_safety import parse_test_postgres_dsn, verify_postgres_test_connection

pytestmark = pytest.mark.postgres
CITIES = ("Тараз", "Шымкент", "Алматы")
USER = 800_000_400


class SyntheticExtractor:
    def __init__(self, *, ambiguous=False):
        self.ambiguous = ambiguous
        self.calls = []

    async def extract(self, text, schema, *, spec_sha256, context_sha256, **kwargs):
        self.calls.append(text)
        assert schema is ProcurementRequestPatch
        assert "Synthetic cable" in text and "USD 17" in text
        request = ParsedProcurementRequest.model_validate(
            {
                "source_type": "document",
                "city": "Тараз",
                "search_scope": "citywide",
                "warnings": ["Unknown tender currency retained in source; no KZT budget."],
                "items": [
                    {
                        "name": "Synthetic cable",
                        "quantity": "2",
                        "unit": "м",
                        "source_reference": "page:1:row:1",
                    },
                    {
                        "name": "Synthetic mat",
                        "quantity": "1",
                        "unit": "шт",
                        "source_reference": "page:1:row:2",
                        "ambiguous_fields": ["thickness"] if self.ambiguous else [],
                    },
                ],
            }
        )
        return schema(spec_sha256=spec_sha256, context_sha256=context_sha256, request=request)


class SyntheticTelegram:
    def __init__(self, content):
        self.content = content
        self.calls = []

    async def get_file(self, file_id):
        self.calls.append(("get_file", file_id))
        return SimpleNamespace(file_path="synthetic-upload")

    async def download_file(self, path, *, destination):
        self.calls.append(("download", path))
        destination.write(self.content)


class SyntheticBrowser:
    """A provider fake; native code validates hashes, evidence and item bindings."""

    def __init__(self):
        self.calls = []

    async def search_suppliers(self, plan):
        self.calls.append(plan)
        candidates = []
        if plan.source == ResearchSource.TWO_GIS:
            external_id = str(700_000_000_000_000 + plan.item_index)
            phone = f"+77000000{400 + plan.item_index:03d}"
            candidates = [
                SupplierCandidate(
                    source=plan.source,
                    external_id=external_id,
                    name=f"Synthetic supplier {plan.item_index}",
                    address="Тараз, Synthetic Street 1",
                    contacts=[
                        SupplierContact(kind=ContactKind.WHATSAPP, value=phone, confidence=1)
                    ],
                    evidence=[
                        SupplierEvidence(
                            evidence_type=EvidenceType.BUSINESS_LISTING,
                            source=plan.source,
                            source_url=f"https://2gis.kz/taraz/firm/{external_id}",
                            observed_at=datetime.now(UTC),
                            raw_snapshot_sha256="b" * 64,
                            claims=[
                                EvidenceClaim.BUSINESS_IDENTITY,
                                EvidenceClaim.ADDRESS,
                                EvidenceClaim.CONTACT,
                            ],
                        )
                    ],
                    contact_discovery_reason="phone_found",
                )
            ]
        return SupplierBrowserResult(
            spec_sha256=BROWSER_SPEC_SHA256,
            request_sha256=_request_hash(plan),
            candidates=candidates,
            exhausted=True,
            pages_examined=1,
        )

    async def resolve_locality(self, task):
        raise AssertionError("citywide fixture must not resolve locality")


class SyntheticWaha:
    session_name = "synthetic-integrated"
    outbound_enabled = True

    def __init__(self):
        self.jid = "77000000999@c.us"
        self.calls = []

    async def get_session_status(self):
        self.calls.append(("status",))
        return "WORKING"

    async def check_phone_exists(self, recipient):
        self.calls.append(("check", recipient))
        return WahaContactCheck(True, recipient.lstrip("+") + "@c.us")

    async def get_account_identity(self):
        self.calls.append(("identity", self.jid))
        return WahaAccountIdentity(self.jid, datetime.now(UTC))

    async def new_message_id(self):
        assert self.outbound_enabled
        self.calls.append(("new_id",))
        return "synthetic-request-1"

    async def send_text(self, recipient, text, **kwargs):
        assert self.outbound_enabled
        self.calls.append(("POST", recipient, text, kwargs))
        return WahaSendResult("synthetic-sent-1", {"id": "synthetic-sent-1"})

    async def fetch_message(self, **kwargs):
        raise AssertionError("a known completed send must not query provider history")


@pytest.fixture
async def connected_pg(tmp_path):
    dsn = os.getenv("TEST_POSTGRES_DSN")
    if not dsn:
        pytest.skip("requires disposable test_uran_integrated PostgreSQL")
    target = parse_test_postgres_dsn(
        dsn,
        expected_database="test_uran_integrated",
        expected_port=55450,
    )
    assert target.host == "127.0.0.1"
    await verify_postgres_test_connection(dsn, target)
    admin = await asyncpg.connect(dsn)
    schema = "uran_connected_" + uuid4().hex
    await admin.execute(f'CREATE SCHEMA "{schema}"')
    pool = None
    try:
        pool = await asyncpg.create_pool(
            dsn,
            min_size=1,
            max_size=6,
            init=_init_connection,
            server_settings={"search_path": schema},
        )
        await run_migrations(pool, Path(__file__).parents[1] / "migrations")
        yield SimpleNamespace(
            pool=pool, root=tmp_path, waha=SyntheticWaha(), browser=SyntheticBrowser()
        )
    finally:
        if pool:
            await pool.close()
        await admin.execute(f'DROP SCHEMA "{schema}" CASCADE')
        await admin.close()


def new_worker(state, *, excluded=()):
    return JobWorker(
        pool=state.pool,
        bot=state.bot,
        storage=state.storage,
        transcriber=None,
        intake=IntakeService(state.extractor),
        worker_id="synthetic-" + uuid4().hex,
        research_executor=ParallelResearchRunExecutor(state.pool, state.browser),
        results_root=state.root / "reports",
        whatsapp_enabled=True,
        waha_session=state.waha.session_name,
        waha_client=state.waha,
        excluded_job_kinds=excluded,
    )


async def drain(state, *kinds):
    """Restart a real worker for every job; unrelated durable work stays queued."""
    for _ in range(150):
        all_kinds = await state.pool.fetch("SELECT DISTINCT kind FROM jobs")
        excluded = tuple(row["kind"] for row in all_kinds if row["kind"] not in kinds)
        if not await new_worker(state, excluded=excluded).run_once():
            break
    else:
        pytest.fail("synthetic native queue did not settle")
    failures = await state.pool.fetch(
        "SELECT kind,status,last_error_code FROM jobs WHERE kind=ANY($1::text[]) "
        "AND status<>'succeeded'",
        list(kinds),
    )
    assert not failures, [dict(row) for row in failures]


async def ingest(state, kind, *, opted_in, table_currency=None):
    path = state.root / f"synthetic.{kind}"
    body = "Synthetic cable 2 m; Synthetic mat 1 piece; source price USD 17"
    if kind == "pdf":
        with fitz.open() as document:
            document.new_page().insert_text((40, 60), body)
            document.save(path)
    else:
        document = Document()
        if table_currency is None:
            document.add_paragraph(body)
        else:
            table = document.add_table(rows=2, cols=7)
            rows = [
                [
                    "№ п/п",
                    "Наименование",
                    "Ед. изм.",
                    "Кол-во",
                    "Цена, " + table_currency,
                    "Стоимость",
                    "Краткая характеристика",
                ],
                ["1", "Synthetic cable", "м", "2", "17", "34", "synthetic source spec"],
            ]
            for row, values in zip(table.rows, rows, strict=True):
                for cell, value in zip(row.cells, values, strict=True):
                    cell.text = value
        document.save(path)
    state.bot = SyntheticTelegram(path.read_bytes())
    state.storage = LocalContentAddressedStorage(state.root / "objects")
    state.extractor = SyntheticExtractor(ambiguous=opted_in)
    # Owner enrollment is an authenticated application operation, before intake.
    owner = await state.pool.fetchval(
        "INSERT INTO users(telegram_id,display_name) VALUES($1,'Synthetic owner') RETURNING id",
        USER,
    )
    state.owner = owner
    if opted_in:
        async with state.pool.acquire() as c:
            await replace_research_profile(
                c,
                owner_id=owner,
                customer_id=None,
                policy=ResearchAutomationPolicy(ordered_cities=CITIES),
                consent_reference="synthetic-owner-action",
                consent_text="Public research only.",
                research_autostart=True,
            )
    kwargs = dict(
        update_id=900_400,
        telegram_user_id=USER,
        chat_id=USER,
        telegram_message_id=1,
        message_kind="document",
        raw_update={"synthetic": True},
        attachment=IncomingAttachment(
            external_file_id="synthetic-file",
            external_file_unique_id="synthetic-unique",
            kind="document",
            original_filename=path.name,
            mime_type="application/pdf"
            if kind == "pdf"
            else "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            declared_size=path.stat().st_size,
        ),
        acknowledgement_text="Synthetic acknowledgement",
    )
    accepted = await persist_telegram_message(state.pool, **kwargs)
    assert accepted and await persist_telegram_message(state.pool, **kwargs) is None
    await drain(state, "download_attachment", "extract_document", "process_intake")
    state.case = await state.pool.fetchval("SELECT id FROM procurement_cases")
    assert state.case
    assert await state.pool.fetchval("SELECT count(*) FROM case_parse_versions") == 1
    assert await state.pool.fetchval("SELECT count(*) FROM source_artifacts") == 1
    source_text = await state.pool.fetchval("SELECT content_text FROM content_extractions")
    assert ("USD 17" if table_currency is None else "Цена, " + table_currency) in source_text
    assert await state.pool.fetchval("SELECT count(*) FROM request_item_financials") == 0
    # Re-enter the actual download/extraction/intake entrypoints on a new worker.
    await new_worker(state)._download_attachment({"attachment_id": str(accepted.attachment_id)})
    extraction = await state.pool.fetchval("SELECT id FROM content_extractions")
    await new_worker(state)._process_intake(
        {
            "incoming_message_id": str(accepted.message_id),
            "extraction_id": str(extraction),
            "input_kind": "document",
        }
    )
    assert state.bot.calls == [("get_file", "synthetic-file"), ("download", "synthetic-upload")]
    assert len(state.extractor.calls) == (1 if table_currency is None else 0)
    return accepted


def signed_headers(user=USER):
    values = {"auth_date": str(int(time.time())), "user": json.dumps({"id": user})}
    check = "\n".join(f"{key}={value}" for key, value in sorted(values.items()))
    secret = hmac.new(b"WebAppData", b"synthetic-token", hashlib.sha256).digest()
    values["hash"] = hmac.new(secret, check.encode(), hashlib.sha256).hexdigest()
    return {"Authorization": "tma " + urlencode(values)}


@pytest.mark.parametrize("kind", ["pdf", "docx"])
@pytest.mark.parametrize("opted_in", [False, True])
async def test_native_document_to_research_outreach_and_authenticated_quote(
    connected_pg,
    kind,
    opted_in,
    monkeypatch,
):
    from procurement_bot.supplier_dialogue import SendPolicy

    # Synthetic clock window only; all send, consent, identity and budget gates run.
    start = (datetime.now(UTC) + timedelta(hours=2)).time().replace(tzinfo=None)
    end = (datetime.now(UTC) + timedelta(hours=3)).time().replace(tzinfo=None)

    async def test_window(*args):
        return SendPolicy(timezone="UTC", quiet_start=start, quiet_end=end)

    for module in ("whatsapp_delivery", "whatsapp_approval_workflow"):
        monkeypatch.setattr(f"procurement_bot.{module}.get_case_send_policy", test_window)
    state = connected_pg
    await ingest(state, kind, opted_in=opted_in)
    await drain(state, "prepare_research")
    approval = await state.pool.fetchrow(
        "SELECT * FROM action_approvals WHERE action_type='start_research'"
    )
    assert approval["status"] == ("consumed" if opted_in else "requested")
    if not opted_in:
        assert not state.browser.calls
        assert await state.pool.fetchval("SELECT count(*) FROM research_policy_authorizations") == 0
        async with state.pool.acquire() as c:
            assert await approve_request(c, approval_id=approval["id"], approved_by=state.owner)
        await start_approved_research(state.pool, approval_id=approval["id"])
    parse = await state.pool.fetchval("SELECT normalized_payload FROM case_parse_versions")
    assert (parse["city"], parse["search_scope"], parse["source_type"]) == (
        "Тараз",
        "citywide",
        "document",
    )
    if opted_in:
        snapshot = await state.pool.fetchrow("SELECT * FROM case_research_policy_snapshots")
        assert snapshot["snapshot"]["policy"]["ordered_cities"] == list(CITIES)
        profile = await state.pool.fetchrow("SELECT * FROM owner_research_automation_profiles")
        assert snapshot["snapshot"]["profile_policy_sha256"] == profile["policy_sha256"]
    replay = await start_approved_research(state.pool, approval_id=approval["id"])
    assert not replay.newly_started
    await drain(state, "execute_research", "parallel_research_task")
    assert await state.pool.fetchval("SELECT count(*) FROM research_runs") == 1
    assert await state.pool.fetchval("SELECT status FROM research_runs") in {"succeeded", "partial"}
    assert state.browser.calls
    assert {call.locality.city for call in state.browser.calls} == {"Тараз"}
    assert {call.item_index for call in state.browser.calls} == ({0} if opted_in else {0, 1})
    assert await state.pool.fetchval("SELECT count(*) FROM offers") == (1 if opted_in else 2)
    # Initial research consent grants no permission to transmit enquiries.
    assert not any(call[0] == "POST" for call in state.waha.calls)
    assert (
        await state.pool.fetchval(
            "SELECT count(*) FROM outbox_events WHERE event_type='whatsapp.send_text'"
        )
        == 0
    )
    # Finish the connected outreach boundary through the real draft/approval code.
    async with state.pool.acquire() as c:
        await ensure_account(c, state.waha.session_name)
    state.waha.outbound_enabled = False
    await drain(state, "prepare_supplier_outreach")
    account = await state.pool.fetchrow("SELECT * FROM whatsapp_account_controls")
    assert account["operator_held"] and not account["pinned_account_jid"]
    assert await state.pool.fetchval("SELECT count(*) FROM supplier_conversations") > 0
    assert not any(call[0] in {"new_id", "POST"} for call in state.waha.calls)
    send = await state.pool.fetchrow(
        "SELECT * FROM action_approvals WHERE action_type='send_whatsapp' "
        "ORDER BY requested_at LIMIT 1"
    )
    assert send
    async with state.pool.acquire() as c:
        assert await approve_request(c, approval_id=send["id"], approved_by=state.owner)
    await start_approved_whatsapp_send(state.pool, approval_id=send["id"])
    event = await state.pool.fetchrow(
        "SELECT * FROM outbox_events WHERE event_type='whatsapp.send_text' LIMIT 1"
    )
    assert event
    sender = WhatsAppOutboxWorker(pool=state.pool, client=state.waha, worker_id="held-send")
    assert await sender.run_once(event_id=event["id"])
    assert (
        await state.pool.fetchval(
            "SELECT last_error_code FROM outbox_events WHERE id=$1", event["id"]
        )
        == "whatsapp_account_operator_hold"
    )
    assert not any(call[0] in {"identity", "new_id", "POST"} for call in state.waha.calls)
    assert await state.pool.fetchval("SELECT count(*) FROM whatsapp_account_send_reservations") == 0
    await pin_whatsapp_account_identity(
        state.pool,
        state.waha,
        actor_user_id=state.owner,
        expected_account_jid=state.waha.jid,
        reason="synthetic exact account",
    )
    pinned = await state.pool.fetchrow("SELECT * FROM whatsapp_account_controls")
    assert pinned["operator_held"] and pinned["pinned_account_jid"] == state.waha.jid
    async with state.pool.acquire() as c:
        await configure_account_budget(
            c,
            state.waha.session_name,
            daily_message_limit=10,
            daily_new_contact_limit=2,
            minimum_interval_seconds=0,
            max_concurrent_posts=1,
            actor_user_id=state.owner,
            reason="synthetic bounded budget",
        )
        await observe_provider_state(
            c,
            state.waha.session_name,
            state="clear",
            observed_at=datetime.now(UTC),
            source="synthetic",
            reason="fake provider healthy",
        )
        await set_account_hold(
            c,
            state.waha.session_name,
            held=False,
            actor_user_id=state.owner,
            reason="synthetic owner release",
        )
    await drain(state, "prepare_supplier_outreach")
    # Wrong live self-account cannot allocate ID, reserve exposure or POST.
    state.waha.outbound_enabled = True
    await state.pool.execute("UPDATE outbox_events SET available_at=now() WHERE id=$1", event["id"])
    state.waha.jid = "77000000888@c.us"
    sender = WhatsAppOutboxWorker(pool=state.pool, client=state.waha, worker_id="synthetic-send")
    assert await sender.run_once(event_id=event["id"])
    assert not any(call[0] in {"new_id", "POST"} for call in state.waha.calls)
    assert (
        await state.pool.fetchval(
            "SELECT provider_request_id FROM outbox_events WHERE id=$1", event["id"]
        )
        is None
    )
    assert await state.pool.fetchval("SELECT count(*) FROM whatsapp_account_send_reservations") == 0
    assert (
        await state.pool.fetchval(
            "SELECT last_error_code FROM outbox_events WHERE id=$1", event["id"]
        )
        == "whatsapp_account_identity_mismatch"
    )
    state.waha.jid = "77000000999@c.us"
    await state.pool.execute("UPDATE outbox_events SET available_at=now() WHERE id=$1", event["id"])
    assert await sender.run_once(event_id=event["id"])
    sent = await state.pool.fetchrow(
        "SELECT status,last_error_code FROM outbox_events WHERE id=$1", event["id"]
    )
    assert sent["status"] == "sent", (dict(sent), state.waha.calls)
    assert not await WhatsAppOutboxWorker(
        pool=state.pool, client=state.waha, worker_id="restart"
    ).run_once(event_id=event["id"])
    assert len([call for call in state.waha.calls if call[0] == "POST"]) == 1
    conversation = event["payload"]["conversation_id"]
    quote = "Есть в наличии. Цена 1250 KZT за м. Доставка в течение 3 дней, самовывоз бесплатно."
    inbound = {
        "event": "message",
        "session": state.waha.session_name,
        "payload": {
            "id": "synthetic-quote-1",
            "from": event["payload"]["recipient"].lstrip("+") + "@c.us",
            "fromMe": False,
            "body": quote,
            "timestamp": int(time.time()),
        },
    }
    ingress = WahaWebhookHttpIngress(
        pool=state.pool,
        shared_secret="synthetic-secret",  # noqa: S106
    )
    body = json.dumps(inbound).encode()
    assert (
        await ingress.handle(
            WahaWebhookHttpRequest("POST", {"Content-Type": "application/json"}, body)
        )
    ).status == 401
    headers = {"Content-Type": "application/json", "X-Webhook-Secret": "synthetic-secret"}
    assert (await ingress.handle(WahaWebhookHttpRequest("POST", headers, body))).status == 200
    assert (await ingress.handle(WahaWebhookHttpRequest("POST", headers, body))).status == 200
    await drain(state, "process_supplier_reply")
    observation = await state.pool.fetchrow(
        "SELECT * FROM offer_observations WHERE source_type='supplier_message'"
    )
    assert observation["price_amount"] == Decimal("1250")
    assert observation["currency"] == "KZT" and observation["price_unit"] == "m"
    assert observation["raw_payload"]["text"] == quote
    async with state.pool.acquire() as c:
        snapshot = await get_current_quote_snapshot(c, observation["offer_id"])
    assert snapshot.facts["price_amount"].value == Decimal("1250")
    assert snapshot.facts["price_unit"].value == "m"
    settings = Settings(
        _env_file=None,
        TELEGRAM_BOT_TOKEN="synthetic-token",  # noqa: S106
        telegram_allowed_user_ids=str(USER),
    )
    async with TestClient(TestServer(create_mini_app(state.pool, settings))) as client:
        path = f"/api/cases/{state.case}"
        assert (await client.get(path)).status == 401
        response = await client.get(path, headers=signed_headers())
        assert response.status == 200
        detail = await response.json()
        if opted_in:
            unresolved = detail["source_visibility"]["items"][1]
            assert unresolved["issues"] == ["thickness"]
            assert unresolved["stage"] == "needs_review"
        matched = next(o for o in detail["offers"] if o["id"] == str(observation["offer_id"]))
        assert Decimal(matched["quote_snapshot_facts"]["price_amount"]["value"]) == 1250
        assert matched["quote_snapshot_facts"]["price_unit"]["value"] == "m"
    assert (
        await state.pool.fetchval(
            "SELECT count(*) FROM supplier_conversation_messages WHERE direction='inbound'"
        )
        == 1
    )
    assert conversation


@pytest.mark.parametrize("currency", ["USD", "не указана"])
async def test_native_docx_currency_guard_without_structured_extraction(connected_pg, currency):
    state = connected_pg
    await ingest(state, "docx", opted_in=True, table_currency=currency)
    parsed = await state.pool.fetchrow("SELECT * FROM case_parse_versions")
    assert parsed["parser_name"] == "priced-table-intake"
    assert parsed["normalized_payload"]["items"][0]["tender_unit_price_private"] is None
    assert parsed["normalized_payload"]["warnings"]
    assert state.extractor.calls == []
    assert await state.pool.fetchval("SELECT count(*) FROM request_item_financials") == 0
