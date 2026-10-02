"""Extract a DOCX tender and build a conservative, auditable valuation JSON.

Usage: PYTHONPATH=src python3 scripts/build_tender_market_valuation.py \
  tender.docx --quotes quotes.json --output valuation.json
"""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import asdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from docx import Document

from procurement_bot.market_valuation import MarketQuote, TenderLine, value_tender


def decimal_text(text: str) -> Decimal:
    return Decimal(re.sub(r"[\s\u00a0\u202f]", "", text).replace(",", "."))


def extract_lines(path: Path) -> tuple[list[TenderLine], Decimal]:
    document = Document(path)
    if len(document.tables) != 1:
        raise ValueError("expected one tender items table")
    rows = document.tables[0].rows
    if len(rows) < 3 or len(rows[0].cells) != 7:
        raise ValueError("unexpected tender table shape")
    lines: list[TenderLine] = []
    for row in rows[1:-1]:
        cells = [cell.text.strip() for cell in row.cells]
        line = TenderLine(
            number=int(cells[0]), name=cells[1], specification=cells[2],
            unit=cells[3], quantity=decimal_text(cells[4]),
            tender_unit_price=decimal_text(cells[5]),
        )
        declared = decimal_text(cells[6])
        if declared != line.tender_total:
            raise ValueError(f"line {line.number}: declared total differs from quantity × price")
        lines.append(line)
    declared_total = decimal_text(rows[-1].cells[-1].text)
    if sum((line.tender_total for line in lines), Decimal(0)) != declared_total:
        raise ValueError("document total does not reconcile")
    return lines, declared_total


def load_quotes(path: Path) -> list[MarketQuote]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, list):
        raise ValueError("quotes file must contain a list")
    quotes: list[MarketQuote] = []
    for row in data:
        values = dict(row)
        for key in (
            "price", "tender_units_per_price_unit", "available_quantity", "order_multiple"
        ):
            if values.get(key) is not None:
                values[key] = Decimal(str(values[key]))
        values["checked_at"] = datetime.fromisoformat(values["checked_at"])
        quotes.append(MarketQuote(**values))
    return quotes


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--quotes", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    lines, declared_total = extract_lines(args.docx)
    quotes = load_quotes(args.quotes)
    unknown = {quote.line_number for quote in quotes} - {line.number for line in lines}
    if unknown:
        raise ValueError(f"quotes reference unknown lines: {sorted(unknown)}")
    valuation = value_tender(lines, quotes)
    payload = {
        "source_document": str(args.docx.resolve()),
        "declared_tender_total": declared_total,
        "generated_at": datetime.now().astimezone(),
        "summary": {
            "line_count": len(lines),
            "covered_count": sum(row.market_total is not None for row in valuation.rows),
            "tender_total": valuation.tender_total,
            "tender_total_covered": valuation.tender_total_covered,
            "market_total_covered": valuation.market_total_covered,
            "difference_covered": valuation.difference_covered,
            "complete_market_total": valuation.complete_market_total,
        },
        "rows": [asdict(row) for row in valuation.rows],
        "quotes": [asdict(quote) for quote in quotes],
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(payload["summary"], ensure_ascii=False, default=str))


if __name__ == "__main__":
    main()
