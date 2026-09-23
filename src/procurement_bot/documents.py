from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import fitz
from docx import Document

from procurement_bot.errors import PermanentProviderError


@dataclass(frozen=True)
class DocumentSegment:
    ordinal: int
    kind: str
    locator: str
    text: str


def extract_pdf(path: Path) -> list[DocumentSegment]:
    try:
        document = fitz.open(path)
    except Exception as exc:
        raise PermanentProviderError("PDF cannot be opened") from exc
    segments: list[DocumentSegment] = []
    try:
        if document.page_count > 500:
            raise PermanentProviderError("PDF exceeds the 500-page safety limit")
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            segments.append(
                DocumentSegment(
                    ordinal=len(segments),
                    kind="pdf_page",
                    locator=f"page:{page_number}",
                    text=text,
                )
            )
    finally:
        document.close()
    return segments


def extract_docx(path: Path) -> list[DocumentSegment]:
    try:
        document = Document(path)
    except Exception as exc:
        raise PermanentProviderError("DOCX cannot be opened") from exc
    segments: list[DocumentSegment] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            segments.append(
                DocumentSegment(len(segments), "docx_paragraph", f"paragraph:{index}", text)
            )
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                segments.append(
                    DocumentSegment(
                        len(segments),
                        "docx_table_row",
                        f"table:{table_index}:row:{row_index}",
                        " | ".join(values),
                    )
                )
    return segments


def extract_document(path: Path, kind: str) -> list[DocumentSegment]:
    if kind == "pdf":
        return extract_pdf(path)
    if kind == "docx":
        return extract_docx(path)
    raise PermanentProviderError(f"unsupported document kind: {kind}")


def joined_text(segments: list[DocumentSegment], *, max_chars: int = 500_000) -> str:
    blocks: list[str] = []
    length = 0
    for segment in segments:
        block = f"[{segment.locator}]\n{segment.text}".strip()
        if not block:
            continue
        if length + len(block) + 2 > max_chars:
            raise PermanentProviderError("document text exceeds intake limit")
        blocks.append(block)
        length += len(block) + 2
    return "\n\n".join(blocks)
